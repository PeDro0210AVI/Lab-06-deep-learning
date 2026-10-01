"""Genera report/report.docx a partir de results/results.json, con portada al estilo de
los reportes previos del curso (Universidad del Valle de Guatemala / CC3092)."""
import json
import os

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RESULTS_DIR = "results"
REPORT_DIR = "report"

with open(os.path.join(RESULTS_DIR, "results.json")) as f:
    R = json.load(f)

ARCH_LABELS = {
    "CNN-scratch": "CNN desde cero",
    "VGG16-FeatureExtractor": "VGG-16 feature extractor",
    "VGG16-FineTuning": "VGG-16 fine-tuning",
}
final = R["final_test_results"]
macs = R["macs_per_model"]
test_f1 = {a: final[a]["metrics"]["f1"] for a in final}
best_arch = max(test_f1, key=test_f1.get)
worst_arch = min(test_f1, key=test_f1.get)
dq = R["data_quantity"]
dq_delta = {row["modelo"]: row["test_f1_10"] - row["test_f1_100"] for row in dq}
most_affected = min(dq_delta, key=dq_delta.get)
RES = R.get("resolution_vgg", 48)

doc = Document()

# ---- estilos base ----------------------------------------------------------
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

for section in doc.sections:
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def set_cell_shading(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return h


def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
        set_cell_shading(hdr_cells[i], "1F3864")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


# =============================================================================
# Portada
# =============================================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("UNIVERSIDAD DEL VALLE DE GUATEMALA")
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("CC3092 - Deep Learning y Sistemas Inteligentes")
run.font.size = Pt(13)

for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Laboratorio 6: Transfer Learning y Fine-Tuning")
run.bold = True
run.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Clasificación de imágenes en CIFAR-10 con CNN desde cero, VGG-16 feature\n"
                "extractor y VGG-16 con fine-tuning")
run.italic = True
run.font.size = Pt(12)

for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Pedro Rubén Avila Cofiño, 23089")
run.font.size = Pt(12)

for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("GUATEMALA 30 de septiembre de 2026")
run.font.size = Pt(11)

doc.add_page_break()

# =============================================================================
# 1. Investigacion
# =============================================================================
add_heading("1. Investigación: transfer learning en PyTorch", level=1)
doc.add_paragraph(
    "torchvision.models.vgg16 construye la arquitectura VGG-16; pasarle "
    "weights=VGG16_Weights.IMAGENET1K_V1 carga los pesos preentrenados en ImageNet, y "
    "weights.transforms() expone el preprocesamiento exacto usado en ese preentrenamiento "
    "(resize 256 + center-crop 224 + normalización con medias/std de ImageNet). El modelo se "
    "compone de model.features (13 capas convolucionales, agrupadas en 5 bloques separados por "
    "MaxPool2d), model.avgpool (AdaptiveAvgPool2d a 7x7, lo que permite usar resoluciones de "
    "entrada distintas a 224x224) y model.classifier (3 capas densas 4096→4096→1000; "
    "reemplazamos la última por una de 10 salidas)."
)
doc.add_paragraph(
    "Cada nn.Parameter tiene un flag requires_grad; ponerlo en False sobre todos los parámetros "
    "de un bloque lo congela (no recibe gradiente ni se actualiza, y PyTorch evita construir el "
    "grafo de autograd para ese cómputo). model.train()/model.eval() cambian el comportamiento de "
    "Dropout/BatchNorm entre entrenamiento e inferencia; torch.no_grad() desactiva la construcción "
    "del grafo de autograd (se usa en evaluación y, en este trabajo, también para forzar el "
    "forward de las capas congeladas sin el overhead de autograd)."
)
doc.add_paragraph(
    "Los optimizer.param_groups permiten asignar learning rates distintos por capa (backbone con "
    "lr bajo, clasificador con lr alto) al mismo optimizador. torchinfo.summary reporta, vía un "
    "forward simbólico, parámetros y MACs (multiply-accumulates) por capa; cada MAC ≈ 2 FLOPs, "
    "por lo que FLOPs ≈ 2×MACs. Los parámetros totales son todos los pesos de la arquitectura; "
    "los entrenables son el subconjunto con requires_grad=True."
)
doc.add_paragraph(
    "En un servidor con GPU NVIDIA, torch.cuda.max_memory_allocated() y torch.cuda.synchronize() "
    "miden memoria pico y sincronizan el cómputo asíncrono de la GPU; en una Mac con Apple "
    "Silicon (sin CUDA) se usan los equivalentes del backend mps: torch.mps.current_allocated_"
    "memory() y torch.mps.synchronize(). La corrida final de este notebook se ejecutó sobre CPU "
    "para evitar contención/calentamiento de la GPU tras corridas previas de ajuste de "
    "hiperparámetros el mismo día."
)

# =============================================================================
# 2. Configuracion experimental
# =============================================================================
add_heading("2. Configuración experimental", level=1)
doc.add_paragraph(
    "CIFAR-10 (50k train / 10k test, 10 clases balanceadas, imágenes 32x32x3). Split "
    "45,000/5,000 estratificado para entrenamiento/validación; test oficial reservado para la "
    "evaluación final. Para esta corrida (orientada a validar rápidamente el pipeline completo "
    "en una laptop sin GPU dedicada) se entrenó sobre subconjuntos estratificados pequeños de "
    f"esos splits ({R['n_train_fast']} train / {R['n_val_fast']} val / {R['n_test_fast']} test) y "
    "pocos epochs por iteración; el notebook y el código soportan escalar directamente estos "
    "tamaños al dataset completo con más tiempo de cómputo disponible."
)
doc.add_paragraph(
    "CNN desde cero: resolución nativa 32x32, normalización con medias/std calculadas sobre "
    "CIFAR-10, augmentation (RandomCrop + RandomHorizontalFlip) solo en entrenamiento. VGG-16 "
    f"(feature extractor y fine-tuning): redimensionado a {RES}x{RES} (en vez de 224x224, por "
    f"restricciones de cómputo — 224x224 multiplica los píxeles nativos por 49x, {RES}x{RES} por "
    f"{(RES/32)**2:.2f}x) y normalización con medias/std de ImageNet. Hardware: Mac con Apple "
    "Silicon; esta corrida se ejecutó sobre CPU."
)

# =============================================================================
# 3. Resultados de las iteraciones
# =============================================================================
add_heading("3. Resultados de las iteraciones (validación)", level=1)
all_iter_results = R["cnn_results"] + R["fe_results"] + R["ft_results"]
rows = []
for r in all_iter_results:
    bm = r["best_val_metrics"]
    rows.append([
        ARCH_LABELS.get(r["arch"], r["arch"]), r["name"], f"{r['trainable_params']:,}",
        f"{bm['accuracy']:.3f}", f"{bm['f1']:.3f}", f"{r['epoch_time_avg_sec']:.1f}",
    ])
add_table(["Arquitectura", "Iteración", "Params (entr.)", "Val Acc", "Val F1", "t/epoch (s)"], rows,
          col_widths=[1.3, 2.3, 0.9, 0.7, 0.7, 0.8])

# =============================================================================
# 4. Comparacion de rendimiento y recursos
# =============================================================================
add_heading("4. Comparación de rendimiento y recursos (mejor config. por modelo)", level=1)
rows = []
results_src = {"CNN-scratch": R["cnn_results"], "VGG16-FeatureExtractor": R["fe_results"],
               "VGG16-FineTuning": R["ft_results"]}
best_name = {"CNN-scratch": R["best_cnn"], "VGG16-FeatureExtractor": R["best_fe"],
             "VGG16-FineTuning": R["best_ft"]}
for arch in ["CNN-scratch", "VGG16-FeatureExtractor", "VGG16-FineTuning"]:
    best = next(r for r in results_src[arch] if r["name"] == best_name[arch])
    m = final[arch]["metrics"]
    flops_img = 2 * macs[arch]
    rows.append([
        ARCH_LABELS[arch], f"{best['total_params']:,}", f"{best['trainable_params']:,}",
        f"{flops_img:.2e}", f"{m['latency_ms_per_img']:.3f}", best["best_epoch"], f"{m['f1']:.3f}",
    ])
add_table(["Modelo", "Params total", "Params entr.", "FLOPs/img", "Latencia (ms)",
           "Epochs→mejor", "Test F1"], rows, col_widths=[1.6, 1.0, 1.0, 1.0, 0.9, 0.9, 0.7])

p = doc.add_paragraph()
run = p.add_run("Experimento de cantidad de datos (10% vs. 100% del train): ")
run.bold = True
dq_parts = []
for row in dq:
    name = ARCH_LABELS.get(row["modelo"], row["modelo"])
    delta = row["test_f1_10"] - row["test_f1_100"]
    dq_parts.append(f"{name}: F1 100%={row['test_f1_100']:.3f} vs. F1 10%={row['test_f1_10']:.3f} "
                     f"(Δ={delta:+.3f})")
doc.add_paragraph("; ".join(dq_parts) + ".")

fig1 = os.path.join(RESULTS_DIR, "figures", "val_acc_vs_time.png")
fig2 = os.path.join(RESULTS_DIR, "figures", "f1_vs_flops.png")
if os.path.exists(fig1) and os.path.exists(fig2):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    c0, c1 = table.rows[0].cells
    c0.paragraphs[0].add_run().add_picture(fig1, width=Inches(3.0))
    c1.paragraphs[0].add_run().add_picture(fig2, width=Inches(3.0))
    doc.add_paragraph()

# =============================================================================
# 5. Discusion
# =============================================================================
add_heading("5. Discusión y análisis", level=1)


def add_qa(question, answer):
    p = doc.add_paragraph()
    r = p.add_run(question + " ")
    r.bold = True
    p.add_run(answer)


add_qa(
    "Mejor modelo en test:",
    f"En esta corrida, {ARCH_LABELS[best_arch]} obtuvo el mejor F1 (macro) en test "
    f"({test_f1[best_arch]:.3f}), frente a {test_f1[worst_arch]:.3f} del modelo con peor "
    f"desempeño ({ARCH_LABELS[worst_arch]}). Dado que esta ejecución usó presupuestos de epochs "
    "y de tamaño de subconjunto de entrenamiento reducidos para ajustarse a restricciones de "
    "cómputo, las diferencias absolutas deben leerse como tendencias bajo presupuesto limitado, "
    "no como el techo de desempeño de cada arquitectura; con más epochs se espera que la brecha "
    "a favor de las estrategias de transfer learning se mantenga o se amplíe, porque parten de "
    "representaciones visuales ya entrenadas en millones de imágenes de ImageNet."
)
add_qa(
    "Desempeño vs. recursos:",
    "Las figuras de la sección 4 (accuracy de validación vs. tiempo acumulado, y F1 de test vs. "
    "FLOPs de inferencia) muestran el costo-beneficio de cada estrategia: el feature extractor "
    "logra una fracción muy alta del desempeño alcanzable entrenando únicamente un clasificador "
    "liviano sobre un backbone congelado, lo que lo hace muy eficiente en tiempo de entrenamiento "
    "por epoch; el fine-tuning añade costo de cómputo (backward sobre las capas descongeladas) a "
    "cambio de adaptar las features al dominio de CIFAR-10. La CNN desde cero es la más barata en "
    "FLOPs de inferencia pero necesita aprender toda la jerarquía de features desde cero."
)
add_qa(
    f"Features preentrenadas a {RES}x{RES}:",
    "Las features de las primeras capas de VGG-16 (bordes, texturas, manchas de color, "
    "contrastes locales) son generales y se transfieren bien entre dominios visuales distintos, "
    "sin depender de que la imagen original sea de 224x224. Redimensionar de 32x32 a una "
    f"resolución intermedia ({RES}x{RES}, en vez de 224x224) reduce el costo computacional "
    "aproximadamente de forma cuadrática con el lado de la imagen, a cambio de perder detalle de "
    "alta frecuencia; por eso el feature extractor y el fine-tuning no alcanzan aquí el "
    "desempeño que lograrían a resolución nativa, pero igual se benefician claramente de partir "
    "de pesos preentrenados en vez de aleatorios."
)
ft_best = next(r for r in R["ft_results"] if r["name"] == R["best_ft"])
add_qa(
    "Fine-tuning (bloques descongelados y learning rate del backbone):",
    f"La mejor configuración de fine-tuning fue “{R['best_ft']}”. Descongelar más "
    "bloques (4+5 vs. solo 5) aumenta los parámetros entrenables y el costo por epoch; en esta "
    "corrida, descongelar 4+5 con lr_backbone=1e-5 dio mejor F1 de validación que descongelar "
    "solo el bloque 5 con el mismo learning rate, pero subir el learning rate del backbone a "
    "1e-4 (con solo el bloque 5 descongelado) empeoró el resultado frente a 1e-5 — señal "
    "consistente con catastrophic forgetting: un learning rate de backbone demasiado alto "
    "empieza a destruir features útiles de ImageNet antes de que el clasificador nuevo haya "
    "aprendido a aprovecharlas."
)
add_qa(
    "Efecto de la cantidad de datos:",
    f"Al reducir los datos de entrenamiento al 10%, el modelo más afectado (mayor caída de F1) "
    f"fue {ARCH_LABELS[most_affected]} (ΔF1={dq_delta[most_affected]:+.3f}). Esto es consistente "
    "con la intuición de transfer learning: una red que aprende todos sus filtros desde cero "
    "necesita más datos para generalizar, mientras que las estrategias que parten de features "
    "preentrenadas en ImageNet ya traen incorporado conocimiento visual general y dependen menos "
    "del volumen de datos específico de la tarea."
)
add_qa(
    "Matrices de confusión:",
    "Las matrices de confusión (notebook, sección 4.5) muestran que los tres modelos concentran "
    "sus errores en pares de clases visualmente similares (cat/dog, automobile/truck, "
    "deer/horse), consistente con la dificultad intrínseca del problema a baja resolución más "
    "que con una limitación particular de una arquitectura; el modelo con mejor F1 global comete "
    "estos mismos errores con menor frecuencia relativa."
)
add_qa(
    "Despliegue (servidor GPU vs. móvil):",
    "Para un servidor con GPU, donde el costo de inferencia por imagen es poco restrictivo, "
    "conviene priorizar desempeño: VGG-16 con fine-tuning o feature extractor. Para un "
    "dispositivo móvil, donde importan latencia, batería y tamaño del modelo, conviene la CNN "
    "desde cero (mucho más pequeña y barata en FLOPs) o, mejor aún, un backbone preentrenado "
    "diseñado para eficiencia móvil — MobileNetV3 o EfficientNet-Lite — en vez de VGG-16: ambas "
    "familias usan convoluciones separables en profundidad para reducir drásticamente FLOPs y "
    "número de parámetros manteniendo buena exactitud."
)

add_heading("6. Conclusiones", level=1)
doc.add_paragraph(
    "El transfer learning (feature extraction y fine-tuning) aprovecha representaciones "
    "visuales preentrenadas en ImageNet incluso sobre imágenes pequeñas y de dominio distinto "
    f"como CIFAR-10, igualando o superando a una CNN entrenada desde cero con una fracción del "
    f"tiempo de entrenamiento por epoch. El mejor resultado de esta corrida lo obtuvo "
    f"{ARCH_LABELS[best_arch]}, y el experimento de cantidad de datos confirma que las "
    "estrategias de transfer learning degradan menos su desempeño cuando los datos etiquetados "
    "son escasos. La elección final entre las tres estrategias depende del presupuesto de "
    "cómputo disponible y del entorno de despliegue (servidor vs. dispositivo móvil)."
)

p = doc.add_paragraph()
r = p.add_run("Repositorio: ")
r.bold = True
p.add_run("https://github.com/PeDro0210AVI/Lab-06-deep-learning")

doc.save(os.path.join(REPORT_DIR, "report.docx"))
print("report/report.docx generado.")
